"""Export artifacts as Markdown, DOCX or PDF (FR-5.11)."""

from __future__ import annotations

import io
import re
from pathlib import Path
from xml.sax.saxutils import escape

from app.chat.service import describe_location
from app.studio.markdown_ir import Block, parse_markdown

FORMATS = {
    "md": ("text/markdown; charset=utf-8", "md"),
    "docx": (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "docx",
    ),
    "pdf": ("application/pdf", "pdf"),
}


def filename(title: str, ext: str) -> str:
    slug = re.sub(r"[^\w\- ]+", "", title).strip().replace(" ", "-") or "export"
    return f"{slug}.{ext}"


def references(citations: list[dict]) -> list[str]:
    lines = []
    for c in sorted(citations, key=lambda c: c["label"]):
        where = describe_location(c.get("location", {}))
        lines.append(f"[{c['label']}] {c['source_name']}{', ' + where if where else ''}")
    return lines


def to_markdown(title: str, content: str, citations: list[dict]) -> bytes:
    text = content.rstrip() + "\n"
    refs = references(citations)
    if refs:
        text += "\n## Sources\n\n" + "\n".join(f"- {line}" for line in refs) + "\n"
    return text.encode("utf-8")


def _with_sources(content: str, citations: list[dict]) -> str:
    refs = references(citations)
    if not refs:
        return content
    return content.rstrip() + "\n\n## Sources\n\n" + "\n".join(f"- {line}" for line in refs)


# --- DOCX ----------------------------------------------------------------------------------


def to_docx(title: str, content: str, citations: list[dict]) -> bytes:
    import docx
    from docx.shared import Inches, Pt

    document = docx.Document()
    document.core_properties.title = title
    for block in parse_markdown(_with_sources(content, citations)):
        if block.kind == "rule":
            document.add_paragraph("―" * 20)
            continue
        if block.kind == "heading":
            paragraph = document.add_heading(level=min(block.level, 4))
        elif block.kind == "bullet":
            style = "List Bullet" if block.level == 1 else f"List Bullet {min(block.level, 3)}"
            paragraph = document.add_paragraph(style=style)
        elif block.kind == "number":
            style = "List Number" if block.level == 1 else f"List Number {min(block.level, 3)}"
            paragraph = document.add_paragraph(style=style)
        elif block.kind == "quote":
            paragraph = document.add_paragraph(style="Quote")
        elif block.kind == "code":
            paragraph = document.add_paragraph()
        else:
            paragraph = document.add_paragraph()
        for run in block.runs:
            r = paragraph.add_run(run.text)
            r.bold = run.bold or None
            r.italic = run.italic or None
            if run.code:
                r.font.name = "Courier New"
                r.font.size = Pt(10)
        if block.kind == "code":
            paragraph.paragraph_format.left_indent = Inches(0.25)
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


# --- PDF -----------------------------------------------------------------------------------

_FONT_CANDIDATES = [
    # (regular, bold, italic, bold-italic)
    (
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf",
        "/usr/share/fonts/truetype/dejavu/DejaVuSans-Oblique.ttf",
        "/usr/share/fonts/truetype/dejavu/DejaVuSans-BoldOblique.ttf",
    ),
    (
        "C:/Windows/Fonts/arial.ttf",
        "C:/Windows/Fonts/arialbd.ttf",
        "C:/Windows/Fonts/ariali.ttf",
        "C:/Windows/Fonts/arialbi.ttf",
    ),
    (
        "/Library/Fonts/Arial Unicode.ttf",
        "/Library/Fonts/Arial Unicode.ttf",
        "/Library/Fonts/Arial Unicode.ttf",
        "/Library/Fonts/Arial Unicode.ttf",
    ),
]


def _register_fonts() -> tuple[str, str]:
    """A system Unicode font family if one exists, else the built-in Helvetica/Courier."""
    from reportlab.lib.fonts import addMapping
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    for files in _FONT_CANDIDATES:
        if all(Path(f).exists() for f in files):
            names = ("LNB", "LNB-Bold", "LNB-Italic", "LNB-BoldItalic")
            for name, path in zip(names, files, strict=True):
                if name not in pdfmetrics.getRegisteredFontNames():
                    pdfmetrics.registerFont(TTFont(name, path))
            addMapping("LNB", 0, 0, names[0])
            addMapping("LNB", 1, 0, names[1])
            addMapping("LNB", 0, 1, names[2])
            addMapping("LNB", 1, 1, names[3])
            return "LNB", "Courier"
    return "Helvetica", "Courier"


def _markup(block: Block) -> str:
    parts = []
    for run in block.runs:
        text = escape(run.text).replace("\n", "<br/>")
        if run.code:
            text = f'<font name="Courier">{text}</font>'
        if run.bold:
            text = f"<b>{text}</b>"
        if run.italic:
            text = f"<i>{text}</i>"
        parts.append(text)
    return "".join(parts) or " "


def to_pdf(title: str, content: str, citations: list[dict]) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import HRFlowable, Paragraph, Preformatted, SimpleDocTemplate, Spacer

    base, mono = _register_fonts()
    body = ParagraphStyle("body", fontName=base, fontSize=10.5, leading=15, spaceAfter=6)
    headings = {
        level: ParagraphStyle(
            f"h{level}",
            parent=body,
            fontName=base,
            fontSize=size,
            leading=size + 4,
            spaceBefore=10,
            spaceAfter=6,
        )
        for level, size in {1: 20, 2: 16, 3: 13, 4: 11.5}.items()
    }
    quote = ParagraphStyle("quote", parent=body, leftIndent=18, textColor="#444444")
    code = ParagraphStyle("code", parent=body, fontName=mono, fontSize=9, leading=12, leftIndent=12)

    story = []
    for block in parse_markdown(_with_sources(content, citations)):
        if block.kind == "heading":
            story.append(Paragraph(f"<b>{_markup(block)}</b>", headings[min(block.level, 4)]))
        elif block.kind in ("bullet", "number"):
            bullet = "•" if block.kind == "bullet" else f"{block.number}."
            style = ParagraphStyle(
                "item",
                parent=body,
                leftIndent=18 * block.level + 6,
                bulletIndent=18 * (block.level - 1),
                spaceAfter=3,
            )
            story.append(Paragraph(_markup(block), style, bulletText=bullet))
        elif block.kind == "quote":
            story.append(Paragraph(f"<i>{_markup(block)}</i>", quote))
        elif block.kind == "code":
            story.append(Preformatted(block.runs[0].text, code))
        elif block.kind == "rule":
            story.append(HRFlowable(width="100%"))
        else:
            story.append(Paragraph(_markup(block), body))
        story.append(Spacer(1, 1))
    out = io.BytesIO()
    SimpleDocTemplate(
        out,
        pagesize=A4,
        title=title,
        leftMargin=2 * cm,
        rightMargin=2 * cm,
        topMargin=2 * cm,
        bottomMargin=2 * cm,
    ).build(story)
    return out.getvalue()


def export(fmt: str, title: str, content: str, citations: list[dict]) -> tuple[bytes, str, str]:
    """(bytes, media type, filename)."""
    media_type, ext = FORMATS[fmt]
    maker = {"md": to_markdown, "docx": to_docx, "pdf": to_pdf}[fmt]
    return maker(title, content, citations), media_type, filename(title, ext)
