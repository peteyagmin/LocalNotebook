import json
import time
from contextlib import ExitStack

import pytest
from fastapi.testclient import TestClient

from app.core.config import RuntimeConfig
from app.llm.fake import FakeLLMClient
from app.main import create_app

TOKEN = "test-token"
ORIGIN = "http://127.0.0.1:8000"
GUIDE_JSON = json.dumps({"summary": "A short summary.", "topics": ["one", "two", "three"]})


@pytest.fixture
def make_client(tmp_path):
    """Build a client (lifespan running) on tmp_path. Calling it twice simulates a restart."""
    stack = ExitStack()

    def _make(fake: FakeLLMClient | None = None, llm_override: str | None = None):
        fake = fake or FakeLLMClient(default_reply=GUIDE_JSON)
        config = RuntimeConfig(
            data_root=tmp_path,
            host="127.0.0.1",
            port=8000,
            llm_base_url_override=llm_override,
            allowed_origins=frozenset({ORIGIN}),
        )
        app = create_app(
            config, token=TOKEN, llm_factory=lambda _s: fake, frontend_dist=tmp_path / "x"
        )
        client = TestClient(app, headers={"Authorization": f"Bearer {TOKEN}"})
        stack.enter_context(client)
        return client

    yield _make
    stack.close()


@pytest.fixture
def client(make_client):
    return make_client()


@pytest.fixture
def configured(client):
    """A client whose models are configured (ingestion requires them)."""
    client.put("/api/v1/settings", json={"chat_model": "chat-x", "embedding_model": "emb-y"})
    return client


@pytest.fixture
def notebook(configured):
    return configured.post("/api/v1/notebooks", json={"name": "NB"}).json()


def wait_idle(client, notebook_id, timeout=20):
    """Block until no source is queued or processing; return the sources."""
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        sources = client.get(f"/api/v1/notebooks/{notebook_id}/sources").json()
        if not any(s["state"] in ("queued", "processing") for s in sources):
            return sources
        time.sleep(0.05)
    raise AssertionError("sources did not settle")


def make_pdf(pages: list[str]) -> bytes:
    """Minimal valid text PDF, one page per string (no external fixtures needed)."""
    objs: list[bytes] = []
    kids = " ".join(f"{3 + 2 * i} 0 R" for i in range(len(pages)))
    objs.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    objs.append(f"<< /Type /Pages /Kids [{kids}] /Count {len(pages)} >>".encode())
    font_id = 3 + 2 * len(pages)
    for i, text in enumerate(pages):
        content_id = 4 + 2 * i
        safe = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        stream = f"BT /F1 12 Tf 72 720 Td ({safe}) Tj ET".encode()
        objs.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents {content_id} 0 R "
            f"/Resources << /Font << /F1 {font_id} 0 R >> >> >>".encode()
        )
        objs.append(f"<< /Length {len(stream)} >>\nstream\n".encode() + stream + b"\nendstream")
    objs.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for n, body in enumerate(objs, start=1):
        offsets.append(len(out))
        out += f"{n} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode()
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    return bytes(out)


def make_docx(path, paragraphs: list[tuple[str, str]], table: list[list[str]] | None = None):
    """paragraphs: (style, text) pairs, e.g. ("Heading 1", "Intro")."""
    import docx

    d = docx.Document()
    for style, text in paragraphs:
        d.add_paragraph(text, style=style)
    if table:
        t = d.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, cell in enumerate(row):
                t.cell(r, c).text = cell
    d.save(str(path))
